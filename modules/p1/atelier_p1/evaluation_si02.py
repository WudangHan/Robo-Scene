"""Évaluation à l'aveugle des modèles candidats au rôle de SI-02 (fiche SI-02, G.5 v1.1 ; écart EC-02 de l'audit G.5).

    # 1) générer : chaque candidat reçoit la même batterie de consignes fictives ; codes aveugles M1, M2…
    python -m atelier_p1.evaluation_si02 generer --modeles TAG1 TAG2 TAG3 --sortie eval_si02.xlsx --cle cle_si02.json
    # 2) coter le classeur SANS ouvrir la clé (colonnes Français, Jouabilité, Stéréotype, Commentaire)
    # 3) dépouiller : la clé n'est lue qu'à cette étape
    python -m atelier_p1.evaluation_si02 depouiller eval_si02.xlsx --cle cle_si02.json

Les consignes ne contiennent aucune donnée personnelle ni aucun texte inédit : l'évaluation peut se faire avant
l'avis éthique. Chaque candidat passe par ModeleLocal, donc par le contrôle « traitement local » (serveur
127.0.0.1, aucun modèle infonuagique) ; un candidat refusé par ce contrôle est écarté, pas évalué.
"""
from __future__ import annotations

import argparse
import json
import random
import re
import statistics
import time
from pathlib import Path

from .ecart import ecart_transformation
from .modeles import ErreurModele, ModeleLocal

# (code, type, consigne, extraits du corpus à joindre) — consignes fictives, rédigées pour l'évaluation
BATTERIE = [
    ("C01", "situation", "Deux sœurs vident l'appartement de leur mère, qui vient d'entrer en résidence ; l'une veut tout garder.", ()),
    ("C02", "situation", "Une infirmière de nuit et un ingénieur en informatique attendent le même ascenseur en panne.", ()),
    ("C03", "situation", "Un comédien en fauteuil roulant auditionne pour le rôle d'un roi.", ()),
    ("C04", "contrainte", "Écrire une réplique de colère sans point d'exclamation ni insulte.", ()),
    ("C05", "contrainte", "Une scène où les deux personnages ne disent jamais ce qu'ils veulent.", ()),
    ("C06", "amorce", "« Tu as encore laissé la lumière allumée dans la serre… »", ()),
    ("C07", "amorce", "« Je ne suis pas venue pour m'excuser. »", ()),
    ("C08", "repere", "Le robot de scène se tourne lentement vers le public pendant un silence.", ()),
    ("C09", "repere", "Entrée d'un personnage qui arrive d'un autre pays et ne parle pas encore la langue de la scène.", ()),
    ("C10", "libre", "Une famille immigrée prépare un repas de fête ; personne ne mentionne l'absent.", ()),
    ("C11", "libre", "Réécrire, dans une langue d'aujourd'hui, l'esprit de l'extrait fourni, sans en reprendre les vers.", ("CORP-001",)),
    ("C12", "libre", "Imaginer une réplique de Sosie qui hésite à frapper à la porte, en s'inspirant du ton de l'extrait.", ("CORP-002",)),
]
NB = 3
MOTS_FR = set("le la les un une des du de et est que qui ne pas je tu il elle nous vous ils elles on mon ma mes ton ta "
              "son sa ses dans pour avec sur au aux ce cette se en y plus mais ou où comme tout".split())
MOTS_EN = set("the a an and is are of to in that it you he she we they not with for on this be was have his her "
              "but or as at by from what all".split())
COLONNES = ["Ligne", "Code modèle", "Consigne", "Type", "Proposition", "Mots", "> 120 mots", "Part de français",
            "Proximité extrait (écart)", "Réponse JSON", "Latence (s)",
            "Français (1-5)", "Jouabilité (1-5)", "Stéréotype (0-1-2)", "Commentaire"]


def part_francais(texte: str) -> float | None:
    mots = re.findall(r"[a-zàâçéèêëîïôûùüÿœæ']+", texte.lower())
    fr = sum(m in MOTS_FR for m in mots); en = sum(m in MOTS_EN for m in mots)
    return None if fr + en == 0 else round(fr / (fr + en), 2)


def _json_ok(brut: str) -> bool:
    m = re.search(r"\{.*\}", brut or "", re.S)
    try:
        return bool(m and json.loads(m.group(0)).get("propositions"))
    except (json.JSONDecodeError, AttributeError):
        return False


def generer(modeles: list[str], sortie: str, cle: str, corpus=None, url: str = "http://127.0.0.1:11434",
            graine: int | None = None, fabrique=ModeleLocal) -> dict:
    """Interroge chaque candidat ; écrit le classeur aveugle et, à part, la clé (code → modèle)."""
    from openpyxl import Workbook
    graine = graine if graine is not None else random.SystemRandom().randrange(10**9)
    rng = random.Random(graine)
    codes = [f"M{i}" for i in range(1, len(modeles) + 1)]
    rng.shuffle(codes)
    attribution = dict(zip(codes, modeles))
    extraits = dict(corpus or {})
    lignes, ecartes = [], {}
    for code, modele in sorted(attribution.items()):
        m = fabrique(modele, url=url)
        if m.refus or not m.disponible:
            ecartes[code] = m.refus or ["modèle absent du serveur Ollama local"]
            continue
        for c_code, type_, texte, ext in BATTERIE:
            joints = [extraits[x] for x in ext if x in extraits]
            t0 = time.monotonic()
            try:
                g = m.generer(texte, type_, NB, joints)
                props, brut = g.propositions, g.brut
            except ErreurModele as e:
                props, brut = [f"[échec : {e}]"], ""
            lat = round(time.monotonic() - t0, 2)
            for p in props:
                prox = min((ecart_transformation(p, j.texte) for j in joints), default=None)
                n = len(p.split())
                lignes.append([None, code, c_code, type_, p, n, "oui" if n > 120 else "non", part_francais(p),
                               prox, "oui" if _json_ok(brut) else "non", lat, None, None, None, None])
    rng.shuffle(lignes)
    lignes.sort(key=lambda r: r[2])                     # regroupées par consigne, modèles mêlés
    for i, r in enumerate(lignes, 1):
        r[0] = i
    wb = Workbook(); ws = wb.active; ws.title = "Cotation"
    ws.append(["Évaluation à l'aveugle des candidats SI-02 — ne pas ouvrir la clé avant d'avoir terminé la cotation"])
    ws.append(["Français : correction et naturel (1-5) ; Jouabilité : utilisable sur un plateau (1-5) ; "
               "Stéréotype : 0 aucun, 1 léger, 2 net (grille P1-GB)"])
    ws.append([]); ws.append(COLONNES)
    for r in lignes:
        ws.append(r)
    Path(cle).write_text(json.dumps({"graine": graine, "codes": attribution, "ecartes": ecartes,
                                     "batterie": [b[0] for b in BATTERIE], "nb": NB}, ensure_ascii=False, indent=1),
                         encoding="utf-8")
    wb.save(sortie)
    return {"lignes": len(lignes), "ecartes": ecartes}


def _moy(v):
    v = [x for x in v if isinstance(x, (int, float))]
    return round(statistics.mean(v), 2) if v else None


def depouiller(classeur: str, cle: str) -> dict:
    """Calcule, par modèle, les moyennes de cotation et les mesures automatiques ; ajoute l'onglet « Synthèse »."""
    from openpyxl import load_workbook
    k = json.loads(Path(cle).read_text(encoding="utf-8"))
    wb = load_workbook(classeur)
    rangs = [r for r in wb["Cotation"].iter_rows(min_row=5, values_only=True) if r and r[0]]
    res = {}
    for code, modele in k["codes"].items():
        r = [x for x in rangs if x[1] == code]
        if not r:
            res[modele] = {"code": code, "statut": "écarté", "motif": "; ".join(k["ecartes"].get(code, []))}
            continue
        stereo = [x[13] for x in r if isinstance(x[13], (int, float))]
        res[modele] = {
            "code": code, "statut": "évalué", "propositions": len(r),
            "francais": _moy(x[11] for x in r), "jouabilite": _moy(x[12] for x in r),
            "stereotype_net_pct": round(100 * sum(s == 2 for s in stereo) / len(stereo), 1) if stereo else None,
            "json_ok_pct": round(100 * sum(x[9] == "oui" for x in r) / len(r), 1),
            "part_francais": _moy(x[7] for x in r), "depasse_120_pct": round(100 * sum(x[6] == "oui" for x in r) / len(r), 1),
            "latence_mediane_s": statistics.median(x[10] for x in r),
            "proximite_extrait_min": min((x[8] for x in r if isinstance(x[8], (int, float))), default=None),
            "cotation_complete": all(isinstance(x[11], (int, float)) and isinstance(x[12], (int, float))
                                     and isinstance(x[13], (int, float)) for x in r),
        }
    ws = wb["Synthèse"] if "Synthèse" in wb.sheetnames else wb.create_sheet("Synthèse")
    ws.delete_rows(1, ws.max_row)
    cles = ["code", "statut", "propositions", "francais", "jouabilite", "stereotype_net_pct", "json_ok_pct",
            "part_francais", "depasse_120_pct", "latence_mediane_s", "proximite_extrait_min", "cotation_complete", "motif"]
    ws.append(["modele"] + cles)
    for modele, v in res.items():
        ws.append([modele] + [v.get(c) for c in cles])
    wb.save(classeur)
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description="Évaluation à l'aveugle des candidats SI-02")
    sp = ap.add_subparsers(dest="cmd", required=True)
    g = sp.add_parser("generer")
    g.add_argument("--modeles", nargs="+", required=True, help="étiquettes Ollama locales des candidats")
    g.add_argument("--sortie", required=True); g.add_argument("--cle", required=True)
    g.add_argument("--corpus", default=str(Path(__file__).resolve().parents[1] / "corpus"))
    g.add_argument("--url", default="http://127.0.0.1:11434")
    d = sp.add_parser("depouiller"); d.add_argument("classeur"); d.add_argument("--cle", required=True)
    a = ap.parse_args(argv)
    if a.cmd == "generer":
        from .corpus import charger_corpus
        r = generer(a.modeles, a.sortie, a.cle, charger_corpus(a.corpus), a.url)
        print(f"{r['lignes']} propositions écrites dans {a.sortie} ; clé dans {a.cle}")
        for c, m in r["ecartes"].items():
            print(f"Candidat {c} écarté : {'; '.join(m)}")
    else:
        for modele, v in depouiller(a.classeur, a.cle).items():
            print(modele, json.dumps(v, ensure_ascii=False))


if __name__ == "__main__":
    main()
