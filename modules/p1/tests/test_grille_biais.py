"""Tests de la grille de biais : alpha de Krippendorff et génération du lot d'essai (mode simulé)."""
import shutil
from pathlib import Path

import pytest

from atelier_p1.grille_biais import alpha_krippendorff, fiabilite, generer

GRILLE = Path(__file__).resolve().parents[2] / "specifications" / "P1-GB_Grille_biais_narratifs_v1.0.xlsx"

# Données de référence ; valeurs attendues calculées avec le paquet indépendant « krippendorff » (PyPI).
REF = [[2, None], [2, 2], [None, None], [0, 0], [None, 2], [None, 0], [0, 2], [1, 0], [0, None], [0, None], [2, 2],
       [None, 0], [None, 0], [None, 0], [0, 0], [0, 0], [None, 0], [2, 1], [2, None], [0, None], [0, 1], [2, 0],
       [0, 2], [1, 2], [None, 0], [1, 1], [0, None], [1, 0], [0, None], [0, 2]]


def test_alpha_valeurs_de_reference():
    assert alpha_krippendorff(REF, "ordinal") == pytest.approx(0.1033, abs=1e-4)
    assert alpha_krippendorff(REF, "nominal") == pytest.approx(0.1031, abs=1e-4)


def test_alpha_accord_parfait_et_cas_limites():
    assert alpha_krippendorff([[0, 0], [1, 1], [2, 2], [1, 1]], "ordinal") == 1.0
    assert alpha_krippendorff([[0, 0], [0, 0]], "nominal") is None          # une seule valeur : alpha indéfini
    assert alpha_krippendorff([[1, None], [None, 2]], "nominal") is None    # aucune unité pairable


@pytest.mark.skipif(not GRILLE.exists(), reason="classeur de la grille absent")
def test_generation_simulee_et_fiabilite(tmp_path):
    from openpyxl import load_workbook
    sortie = tmp_path / "essai.xlsx"
    assert generer(str(GRILLE), str(sortie), simule=True) == 50
    wb = load_workbook(sortie)
    ws, ws2 = wb["Cotation"], wb["Recotation"]
    for r in range(5, 55):
        for c in range(5, 13):
            ws.cell(row=r, column=c, value=(r + c) % 3)
            ws2.cell(row=r, column=c - 3, value=(r + c) % 3)
    wb.save(sortie)
    res = fiabilite(str(sortie))
    assert res["B1"]["alpha_ordinal"] == 1.0 and res["B1"]["paires"] == 50
