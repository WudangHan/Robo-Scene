"""Tests du contrôle « traitement local » de SI-02 et de l'évaluation à l'aveugle (G.5 v1.1, EC-02)."""
import json
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RACINE))
from atelier_p1.corpus import charger_corpus  # noqa: E402
from atelier_p1.evaluation_si02 import BATTERIE, NB, depouiller, generer, part_francais  # noqa: E402
from atelier_p1.modeles import ErreurModele, ModeleLocal, verifier_local  # noqa: E402


class FauxOllama(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        corps = json.dumps({"models": [{"name": "cand-a:latest"}, {"name": "cand-b:latest"},
                                       {"name": "piege:latest", "remote_host": "https://ollama.com:443"}]}).encode()
        self.send_response(200); self.end_headers(); self.wfile.write(corps)

    def do_POST(self):
        n = int(self.headers["Content-Length"]); corps = json.loads(self.rfile.read(n))
        texte = ('{"propositions": ["La sœur aînée pose le carton et se tait.", "Elle rit dans le noir de la serre.", '
                 '"Le roi entre par la cour."]}') if corps["model"] == "cand-a" else "1. The king enters\n2. Le roi sort"
        rep = json.dumps({"model": corps["model"], "message": {"content": texte}, "eval_count": 30}).encode()
        self.send_response(200); self.send_header("Content-Length", str(len(rep))); self.end_headers(); self.wfile.write(rep)


@pytest.fixture
def ollama():
    h = ThreadingHTTPServer(("127.0.0.1", 0), FauxOllama)
    threading.Thread(target=h.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{h.server_address[1]}"
    h.shutdown()


def test_controle_traitement_local():
    assert verifier_local("apertus:8b", "http://127.0.0.1:11434") == []
    assert verifier_local("apertus:8b", "http://localhost:11434") == []
    assert verifier_local("gpt-oss:120b-cloud", "http://127.0.0.1:11434")          # modèle infonuagique
    assert verifier_local("apertus:8b", "http://192.168.0.12:11434")               # serveur d'une autre machine
    m = ModeleLocal("qwen3-coder:480b-cloud")
    assert not m.disponible
    with pytest.raises(ErreurModele, match="non local"):
        m.generer("Une scène", "libre", 2)


def test_modele_annonce_distant_refuse(ollama):
    m = ModeleLocal("piege", url=ollama, api="ollama")
    assert not m.disponible and "distant" in m.refus[0] and "Ollama" in m.refus[0]
    assert ModeleLocal("cand-a", url=ollama, api="ollama").disponible


def test_generation_aveugle_et_depouillement(ollama, tmp_path):
    sortie, cle = tmp_path / "eval.xlsx", tmp_path / "cle.json"
    r = generer(["cand-a", "cand-b", "piege", "absent"], str(sortie), str(cle), charger_corpus(RACINE / "corpus"),
                url=ollama, graine=7, api="ollama")
    k = json.loads(cle.read_text(encoding="utf-8"))
    assert set(k["codes"].values()) == {"cand-a", "cand-b", "piege", "absent"} and len(k["ecartes"]) == 2
    rangs = [x for x in load_workbook(sortie)["Cotation"].iter_rows(min_row=5, values_only=True) if x[0]]
    assert len(rangs) == len(BATTERIE) * NB + len(BATTERIE) * 2         # cand-a : 3 par consigne ; cand-b : 2
    assert not any("cand-" in str(x) for row in rangs for x in row)     # aucun nom de modèle dans le classeur
    assert [x[2] for x in rangs] == sorted(x[2] for x in rangs)
    wb = load_workbook(sortie); ws = wb["Cotation"]
    for row in ws.iter_rows(min_row=5):
        if row[0].value:
            row[11].value, row[12].value, row[13].value = 4, 3, 0
    wb.save(sortie)
    res = depouiller(str(sortie), str(cle))
    assert res["cand-a"]["json_ok_pct"] == 100.0 and res["cand-b"]["json_ok_pct"] == 0.0
    assert res["cand-a"]["francais"] == 4 and res["cand-a"]["cotation_complete"]
    assert res["piege"]["statut"] == "écarté" and "Synthèse" in load_workbook(sortie).sheetnames
    assert part_francais("Le roi entre dans la salle et se tait") == 1.0 and part_francais("The king is in the hall") == 0.0
